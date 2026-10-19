import contextlib
import io
import os
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import pandas as pd
from openpyxl import load_workbook

import aluminum_cleaner as app


def source_frame():
    # 10 six-metre bars weighing 15 kg => 0.25 kg/metre.
    return pd.DataFrame([
        {'פרויקט': 'TEST', 'שם-מקייט': '02008', 'תאור': 'Profile',
         'גוון': 'WHITE', 'מידה': '6000', 'כמות': '10', 'משקל': '15', 'ספק': 'A'},
        {'פרויקט': 'TEST', 'שם-מקייט': '02008', 'תאור': 'Profile',
         'גוון': 'BLACK', 'מידה': '3000', 'כמות': '2', 'משקל': '1.5', 'ספק': 'B'},
    ])


class MaterialOrderTests(unittest.TestCase):
    def test_pavel_example_and_changed_length(self):
        original = source_frame()
        rows = app.material_order_rows(original)
        result = app.apply_material_order(original, rows, [('14', '6'), ('14', '3')])
        self.assertEqual(result['משקל'].tolist(), [21, 10.5])
        self.assertEqual(result['מידה'].tolist(), [6000, 3000])
        self.assertEqual(result['כמות'].tolist(), [14, 14])
        self.assertEqual(original['משקל'].tolist(), ['15', '1.5'])

    def test_duplicate_sku_colors_and_nonconsecutive_indices_stay_separate(self):
        frame = source_frame().set_axis([4, 19])
        rows = app.material_order_rows(frame)
        result = app.apply_material_order(frame, rows, [(1, 2), (3, 4)])
        self.assertEqual(result['משקל'].tolist(), [0.5, 3])
        self.assertEqual(result['גוון'].tolist(), ['WHITE', 'BLACK'])
        self.assertEqual(result['ספק'].tolist(), ['A', 'B'])

    def test_decimal_comma_and_zero_quantity(self):
        self.assertEqual(app._material_values('14', '6,5', 0.25), (14, 6.5, 22.75))
        self.assertEqual(app._material_values('0', '6', 0.25), (0, 6, 0))

    def test_invalid_edits(self):
        for quantity, length in [('-1', '6'), ('1.5', '6'), ('', '6'), ('nan', '6'),
                                 ('1', '0'), ('1', '-6'), ('1', 'abc'), ('1', 'inf'),
                                 ('1e308', '1e308')]:
            with self.subTest(quantity=quantity, length=length), self.assertRaises(ValueError):
                app._material_values(quantity, length, 0.25)

    def test_missing_or_invalid_original_weight_fails_explicitly(self):
        for value in [None, 'NaN', '0', '-1', 'inf']:
            frame = source_frame()
            frame.loc[0, 'משקל'] = value
            with self.subTest(value=value), self.assertRaisesRegex(ValueError, '02008'):
                app.material_order_rows(frame)
        with self.assertRaisesRegex(ValueError, 'משקל'):
            app.material_order_rows(source_frame().drop(columns='משקל'))

    def test_export_filters_shutters_and_zero_count_and_writes_numeric_rtl_excel(self):
        with tempfile.TemporaryDirectory() as folder:
            inp, out = Path(folder) / 'input.csv', Path(folder) / 'order.xlsx'
            frame = source_frame()
            shutter = frame.iloc[0].copy()
            shutter['שם-מקייט'] = '13090002'
            pd.concat([frame, shutter.to_frame().T], ignore_index=True).to_csv(inp, index=False, sep=';', encoding='utf-8-sig')
            with patch.object(sys, 'argv', ['aluminum_cleaner.py']), \
                 patch.object(app, 'TK_AVAILABLE', True), \
                 patch.object(app, 'pick_files_interactively', return_value=(str(inp), str(out), None)), \
                 patch.object(app, 'input_material_order', return_value=[(14, 6), (0, 3)]) as dialog, \
                 contextlib.redirect_stdout(io.StringIO()):
                app.main()
            self.assertEqual(len(dialog.call_args.args[0]), 2)
            workbook = load_workbook(out)
            sheet = workbook.active
            self.assertTrue(sheet.sheet_view.rightToLeft)
            self.assertEqual(sheet.max_row, 2)
            cells = dict(zip([c.value for c in sheet[1]], [c.value for c in sheet[2]]))
            self.assertEqual(cells['שם-מקייט'], '02008')
            self.assertEqual(cells['כמות'], 14)
            self.assertEqual(cells['מידה'], 6000)
            self.assertEqual(cells['משקל'], 21)
            workbook.close()

    def test_cancel_does_not_create_or_overwrite_output(self):
        with tempfile.TemporaryDirectory() as folder:
            inp, out = Path(folder) / 'input.xlsx', Path(folder) / 'existing.xlsx'
            source_frame().to_excel(inp, index=False)
            out.write_bytes(b'keep existing file')
            with patch.object(sys, 'argv', ['aluminum_cleaner.py']), \
                 patch.object(app, 'TK_AVAILABLE', True), \
                 patch.object(app, 'pick_files_interactively', return_value=(str(inp), str(out), None)), \
                 patch.object(app, 'input_material_order', return_value=None):
                app.main()
            self.assertEqual(out.read_bytes(), b'keep existing file')


@unittest.skipUnless(app.TK_AVAILABLE, 'Tk is unavailable')
class MaterialDialogTests(unittest.TestCase):
    def run_dialog(self, action):
        real_tk = app.tk.Tk
        errors = []

        def create_root():
            root = real_tk()
            def run_action():
                try:
                    action(root)
                except Exception as exc:
                    errors.append(exc)
                    root.destroy()
            root.after(150, run_action)
            root.after(5000, root.destroy)
            return root

        with patch.object(app.tk, 'Tk', side_effect=create_root):
            result = app.input_material_order(app.material_order_rows(source_frame()))
        if errors:
            raise errors[0]
        return result

    def widgets(self, root):
        for widget in root.winfo_children():
            yield widget
            yield from self.widgets(widget)

    def test_real_widgets_edit_and_submit_without_tcl_error(self):
        def action(root):
            entries = [w for w in self.widgets(root) if isinstance(w, app.ttk.Entry)]
            # Entries are created length, quantity for each source row.
            entries[0].delete(0, 'end'); entries[0].insert(0, '6,5')
            entries[1].delete(0, 'end'); entries[1].insert(0, '14')
            labels = [w.cget('text') for w in self.widgets(root) if isinstance(w, app.ttk.Label)]
            self.assertIn('22.750', labels)
            self.assertIn('02008 (WHITE)', labels)
            button = next(w for w in self.widgets(root) if isinstance(w, app.ttk.Button) and w.cget('text') == 'אישור והפקת אקסל')
            button.invoke()
        self.assertEqual(self.run_dialog(action), [(14, 6.5), (2, 3)])

    def test_real_widgets_validation_keeps_dialog_open_and_cancel_returns_none(self):
        def action(root):
            entry = next(w for w in self.widgets(root) if isinstance(w, app.ttk.Entry))
            entry.delete(0, 'end'); entry.insert(0, '-1')
            button = next(w for w in self.widgets(root) if isinstance(w, app.ttk.Button) and w.cget('text') == 'אישור והפקת אקסל')
            with patch.object(app.messagebox, 'showerror') as error:
                button.invoke()
                error.assert_called_once()
            self.assertTrue(root.winfo_exists())
            root.destroy()
        self.assertIsNone(self.run_dialog(action))

    @unittest.skipUnless(sys.platform == 'win32', 'Windows screenshot')
    def test_dialog_screenshot_and_long_list_scrolling(self):
        real_tk = app.tk.Tk
        errors = []

        def create_root():
            root = real_tk()
            def capture():
                try:
                    root.update_idletasks()
                    canvases = [w for w in self.widgets(root) if isinstance(w, app.tk.Canvas)]
                    canvas = canvases[0]
                    canvas.yview_moveto(1)
                    root.update_idletasks()
                    self.assertGreater(canvas.yview()[0], 0)
                    buttons = [w for w in self.widgets(root) if isinstance(w, app.ttk.Button)]
                    self.assertTrue(all(w.winfo_ismapped() for w in buttons))
                    canvas.yview_moveto(0)
                    root.update()
                    screenshot = os.environ.get('KOSTIKA_UI_SCREENSHOT')
                    if screenshot:
                        # Capture only this test window on the Windows CI desktop.
                        script = (
                            'Add-Type -AssemblyName System.Drawing; '
                            f'$bitmap = New-Object System.Drawing.Bitmap({root.winfo_width()}, {root.winfo_height()}); '
                            '$graphics = [System.Drawing.Graphics]::FromImage($bitmap); '
                            f'$graphics.CopyFromScreen({root.winfo_rootx()}, {root.winfo_rooty()}, 0, 0, $bitmap.Size); '
                            "$bitmap.Save($env:KOSTIKA_UI_SCREENSHOT, [System.Drawing.Imaging.ImageFormat]::Png); "
                            '$graphics.Dispose(); $bitmap.Dispose()'
                        )
                        subprocess.run(['powershell', '-NoProfile', '-Command', script], check=True)
                except Exception as exc:
                    errors.append(exc)
                finally:
                    root.destroy()
            root.after(500, capture)
            return root

        rows = app.material_order_rows(source_frame())
        many_rows = [dict(rows[i % 2], sku=str(2000 + i)) for i in range(70)]
        with patch.object(app.tk, 'Tk', side_effect=create_root):
            self.assertIsNone(app.input_material_order(many_rows))
        if errors:
            raise errors[0]


if __name__ == '__main__':
    unittest.main()
